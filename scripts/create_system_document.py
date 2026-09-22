from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "documentacion"
OUTPUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUTPUT_DIR / "Manual_funcional_AP9_Carbon_Twin.docx"

NAVY = "15324A"
TEAL = "0F766E"
LIGHT_BLUE = "EAF3F7"
PALE = "F5F8FA"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(76, 86, 96)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    table = doc.add_table(rows=len(items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)
    for index, item in enumerate(items, 1):
        number_cell, text_cell = table.rows[index - 1].cells
        number_cell.width = Inches(0.2)
        text_cell.width = Inches(6.55)
        set_cell_margins(number_cell, 25, 0, 45, 45)
        set_cell_margins(text_cell, 25, 0, 45, 0)
        number_cell.text = f"{index}."
        text_cell.text = item
        text_cell.paragraphs[0].add_run().add_break()
        number_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number_cell.paragraphs[0].runs[0].bold = True
        for cell in (number_cell, text_cell):
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.08
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, header in enumerate(headers):
        cell = hdr.cells[index]
        cell.text = header
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if col_index > 0 and len(str(value)) < 25:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = GRAY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.13

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(30)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(14)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle_style = styles["Subtitle"]
subtitle_style.font.name = "Aptos Display"
subtitle_style.font.color.rgb = BLACK
subtitle_style.font.italic = False

for style_name, size, before, after in (
    ("Heading 1", 18, 15, 8),
    ("Heading 2", 14, 12, 5),
    ("Heading 3", 11.5, 9, 4),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = False

for style_name in ("List Bullet", "List Number"):
    style = styles[style_name]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)
    style.font.color.rgb = GRAY
    style.paragraph_format.space_after = Pt(3)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "AP 9 Carbon Twin   Documento funcional   "
footer_p.style = styles["Normal"]
footer_p.runs[0].font.size = Pt(8)
footer_p.runs[0].font.color.rgb = RGBColor(100, 110, 120)
add_page_number(footer_p)

# Portada
doc.add_paragraph("AP 9", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.LEFT
p = doc.add_paragraph(style="Title")
p.add_run("Manual funcional del sistema Carbon Twin")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(28)
r = subtitle.add_run("Análisis del coste de carbono de gemelos digitales agrícolas")
r.font.name = "Aptos Display"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = BLACK

add_body(doc, "Este documento explica qué hace el sistema, cómo se relacionan sus componentes y cómo se utiliza para evaluar la huella ambiental de una explotación agrícola digitalizada. Su objetivo es servir como guía funcional y técnica para usuarios, docentes, evaluadores y desarrolladores.")
add_body(doc, "Conclusión principal: Carbon Twin integra inventario tecnológico, evaluación de ciclo de vida, simulación de escenarios, visualización tridimensional, inteligencia artificial y aprendizaje automático en una aplicación web conectada a FastAPI y PostgreSQL.", "Conclusión principal:")

doc.add_paragraph().paragraph_format.space_after = Pt(60)
meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
set_table_borders(meta)
metadata = [
    ("Tipo de documento", "Manual funcional y técnico"),
    ("Sistema", "AP 9 Carbon Twin"),
    ("Tecnologías", "Next.js 15, React 19, FastAPI y PostgreSQL 18"),
    ("Fecha", "Septiembre de 2026"),
]
for i, (label, value) in enumerate(metadata):
    meta.cell(i, 0).text = label
    meta.cell(i, 1).text = value
    set_cell_fill(meta.cell(i, 0), LIGHT_BLUE)
    for cell in meta.rows[i].cells:
        set_cell_margins(cell, 130, 150, 130, 150)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(10)
    meta.cell(i, 0).paragraphs[0].runs[0].bold = True

doc.add_page_break()

doc.add_heading("Contenido", level=1)
contents = [
    "1 Visión general del sistema",
    "2 Arquitectura y flujo de información",
    "3 Módulos funcionales",
    "4 Gemelos digitales agrícolas en 3D",
    "5 Metodología de cálculo de carbono",
    "6 Inteligencia artificial y aprendizaje automático",
    "7 Persistencia seguridad y trazabilidad",
    "8 Flujo de uso recomendado",
    "9 Puesta en marcha sin Docker",
    "10 Alcance actual y conclusiones",
]
for item in contents:
    p = doc.add_paragraph(item)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(7)
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.bold = True

doc.add_page_break()

doc.add_heading("1 Visión general del sistema", level=1)
add_body(doc, "Carbon Twin es una plataforma para estudiar el impacto ambiental asociado a la digitalización de cultivos. El sistema registra la infraestructura física y digital de un proyecto agrícola, estima las emisiones de gases de efecto invernadero durante su ciclo de vida y presenta los resultados mediante paneles, comparaciones y documentos exportables.")
add_body(doc, "La unidad central de resultados es el kilogramo de dióxido de carbono equivalente, abreviado como kg CO2e. Esta unidad permite expresar con una sola medida el efecto climático de diferentes gases y actividades.")

doc.add_heading("1.1 Problema que resuelve", level=2)
add_body(doc, "La agricultura de precisión emplea sensores, gateways, computación Edge, bases de datos, servicios Cloud y modelos digitales. Estas tecnologías mejoran la eficiencia productiva, pero también consumen energía y requieren fabricar, transportar, mantener y retirar equipos. Carbon Twin permite hacer visible ese coste ambiental y compararlo con alternativas tecnológicas.")

doc.add_heading("1.2 Capacidades principales", level=2)
add_bullets(doc, [
    "Administrar proyectos agrícolas y su periodo de evaluación.",
    "Inventariar sensores IoT, dispositivos Edge y recursos Cloud.",
    "Crear y configurar gemelos digitales vinculados a cada proyecto.",
    "Calcular la huella por componente, categoría y etapa del ciclo de vida.",
    "Comparar arquitecturas centralizadas, Edge e híbridas.",
    "Visualizar el cultivo y la telemetría en un entorno tridimensional interactivo.",
    "Generar recomendaciones mediante un asesor de inteligencia artificial.",
    "Entrenar y evaluar modelos de aprendizaje automático para estimación de emisiones.",
    "Exportar resultados a PDF, Word y Excel y conservar un registro de auditoría.",
])

doc.add_heading("1.3 Usuarios previstos", level=2)
add_table(doc, ["Perfil", "Uso principal", "Nivel de acceso"], [
    ["Administrador", "Configura inventarios, usuarios, escenarios y parámetros", "Completo"],
    ["Investigador o analista", "Evalúa resultados, ajusta modelos y genera informes", "Edición analítica"],
    ["Consulta", "Revisa paneles, indicadores e informes", "Solo lectura"],
], [1.5, 3.7, 1.5])

doc.add_page_break()
doc.add_heading("2 Arquitectura y flujo de información", level=1)
add_body(doc, "La solución utiliza una arquitectura de tres capas. El navegador presenta la interfaz y realiza los cálculos interactivos; FastAPI expone los servicios de persistencia, inteligencia artificial y aprendizaje automático; PostgreSQL conserva los recursos del sistema.")

add_table(doc, ["Capa", "Tecnología", "Responsabilidad"], [
    ["Presentación", "Next.js 15, React 19, TypeScript, Tailwind CSS", "Navegación, formularios, paneles, gráficos y exportaciones"],
    ["Visualización", "Three.js y Recharts", "Escena 3D del cultivo y gráficos de indicadores"],
    ["Servicios", "Python FastAPI y SQLAlchemy", "API CRUD, asesor IA, entrenamiento ML y predicción"],
    ["Datos", "PostgreSQL 18", "Persistencia de proyectos, inventarios, escenarios, usuarios y auditorías"],
], [1.25, 2.25, 3.2])

doc.add_heading("2.1 Flujo principal", level=2)
add_numbered(doc, [
    "El usuario selecciona un proyecto agrícola en la cabecera.",
    "La aplicación obtiene de PostgreSQL el inventario asociado mediante FastAPI.",
    "El motor LCA combina cantidades, consumos, horas de operación y factores de emisión.",
    "El dashboard muestra totales, distribución tecnológica, etapas y evolución temporal.",
    "Los módulos de escenarios, IA y ML utilizan el mismo contexto para analizar alternativas.",
    "Las exportaciones convierten los resultados en documentos aptos para revisión y auditoría.",
])

doc.add_heading("2.2 Continuidad sin conexión temporal", level=2)
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(0)
spacer.paragraph_format.line_spacing = Pt(1)
add_body(doc, "El navegador mantiene una copia de trabajo en almacenamiento local. Al recuperar la conexión, la aplicación fusiona los registros locales y remotos por identificador y completa en PostgreSQL los elementos ausentes sin eliminar los ya guardados. Esta estrategia evita que una carga remota parcial reduzca el inventario visible.")

doc.add_page_break()
doc.add_heading("3 Módulos funcionales", level=1)
add_body(doc, "La interfaz organiza las funciones en gestión, infraestructura, cálculo, sostenibilidad y gobierno del sistema. La siguiente matriz resume el propósito de cada módulo.")
modules = [
    ["Dashboard principal", "Resume huella total, energía, infraestructura, categorías, etapas y escenarios."],
    ["Proyectos agrícolas", "Crea y edita explotaciones, superficie, cultivo, ubicación y horizonte de análisis."],
    ["Sensores IoT", "Registra tipos, cantidades, potencia, materiales, comunicaciones y vida útil."],
    ["Infraestructura Edge", "Gestiona gateways y nodos de procesamiento en campo, incluida alimentación solar."],
    ["Recursos Cloud y BD", "Modela cómputo, almacenamiento, región, PUE, energía renovable y factor eléctrico."],
    ["Gemelos digitales", "Define modelos, frecuencia de actualización, volumen de datos y horas de ejecución."],
    ["Ciclo de vida", "Presenta emisiones por fabricación, transporte, instalación, operación, mantenimiento y fin de vida."],
    ["Factores de emisión", "Administra factores técnicos, unidades, fuentes y valores predeterminados."],
    ["Comparar escenarios", "Contrasta arquitectura Cloud, Edge e híbrida en carbono, energía, coste y latencia."],
    ["Informes", "Genera PDF ejecutivo, Word editable y Excel con hojas de cálculo estructuradas."],
    ["Asistente IA", "Produce diagnósticos y recomendaciones contextualizadas para reducir emisiones."],
    ["Pipeline ML", "Entrena cinco modelos, aplica validación y pruebas estadísticas y expone predicciones."],
    ["Usuarios y RBAC", "Gestiona usuarios, roles, organizaciones y estado de acceso."],
    ["Registro de auditoría", "Conserva acciones de creación, modificación, cálculo, exportación y eliminación."],
]
add_table(doc, ["Módulo", "Función"], modules, [1.8, 5.0])

doc.add_page_break()
doc.add_heading("4 Gemelos digitales agrícolas en 3D", level=1)
add_body(doc, "El módulo de gemelos digitales representa de forma visual la parcela y su infraestructura de observación. No es una imagen estática: utiliza una escena WebGL con controles de cámara, iluminación, capas y animaciones que ayudan a interpretar cómo circulan los datos desde el campo hasta el modelo digital.")

doc.add_heading("4.1 Elementos representados", level=2)
add_bullets(doc, [
    "Bloque de suelo, surcos y distribución espacial del cultivo.",
    "Tuberías, líneas de riego y emisores asociados al balance hídrico.",
    "Sensores IoT con indicadores de estado, paneles solares y cobertura inalámbrica.",
    "Gateway Edge que concentra y procesa la telemetría de campo.",
    "Núcleo holográfico que simboliza el modelo digital y sus algoritmos.",
    "Paquetes animados que muestran el flujo de datos entre sensores y núcleo del gemelo.",
    "Zonas NDVI para identificar diferencias de vigor vegetal.",
])

doc.add_heading("4.2 Controles disponibles", level=2)
add_table(doc, ["Control", "Resultado"], [
    ["Órbita", "Permite rotar libremente alrededor de la parcela y acercar o alejar la cámara."],
    ["Satelital", "Coloca la cámara sobre el cultivo para una lectura espacial de sectores y NDVI."],
    ["A ras de campo", "Sitúa la cámara cerca del cultivo para observar sensores y líneas de riego."],
    ["Capas operativas", "Activa o desactiva NDVI, IoT, flujo de datos y red de riego."],
    ["Hora de simulación", "Modifica posición e intensidad solar, ambiente, fondo y niebla."],
    ["Pausa", "Detiene los paquetes, pulsos e indicadores sin perder la vista actual."],
    ["Ampliar", "Expande el visor para inspeccionar la escena con mayor detalle."],
], [1.55, 5.25])

doc.add_heading("4.3 Indicadores del visor", level=2)
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(0)
spacer.paragraph_format.line_spacing = Pt(1)
add_body(doc, "El panel muestra fidelidad, confianza, latencia y carbono del gemelo. La fidelidad se relaciona con la frecuencia de actualización, la confianza con el volumen de telemetría, la latencia con el intervalo de refresco y el carbono con la sobrecarga anual declarada o estimada del cómputo.")

doc.add_page_break()
doc.add_heading("5 Metodología de cálculo de carbono", level=1)
add_body(doc, "El motor aplica una evaluación de ciclo de vida alineada con la estructura de ISO 14040 e ISO 14044. El resultado total agrega emisiones incorporadas y operativas y descuenta, cuando corresponde, créditos de reciclaje o recuperación al final de la vida útil.")

doc.add_heading("5.1 Etapas evaluadas", level=2)
add_table(doc, ["Etapa", "Qué representa"], [
    ["Fabricación", "Materiales, componentes electrónicos, servidores y equipos necesarios."],
    ["Transporte", "Traslado de dispositivos hasta la explotación agrícola."],
    ["Instalación", "Despliegue, calibración inicial y puesta en servicio."],
    ["Operación", "Electricidad consumida durante el periodo de análisis."],
    ["Mantenimiento", "Repuestos, desplazamientos, actualizaciones y recalibraciones."],
    ["Fin de vida", "Gestión RAEE, reciclaje y créditos por recuperación de materiales."],
], [1.45, 5.35])

doc.add_heading("5.2 Estructura del cálculo", level=2)
add_body(doc, "Huella total = fabricación + transporte + instalación + operación + mantenimiento + fin de vida.")
add_body(doc, "En la fase operativa se combinan potencia, horas de uso, factor de emisión de la red y, para recursos Cloud, eficiencia energética del centro de datos. El sistema calcula además emisiones por hectárea, emisiones medias diarias, energía total y participación porcentual de cada categoría.")

doc.add_heading("5.3 Categorías tecnológicas", level=2)
add_bullets(doc, [
    "Sensores IoT y sus comunicaciones.",
    "Dispositivos Edge de procesamiento local.",
    "Infraestructura Cloud, almacenamiento y bases de datos.",
    "Gemelos digitales y entrenamiento o recalibración de modelos.",
    "Red y comunicaciones dedicadas.",
])

doc.add_heading("5.4 Comparación de escenarios", level=2)
add_body(doc, "El simulador permite cambiar número de sensores, nodos Edge, horas Cloud, almacenamiento, frecuencia de transmisión, uso de energía renovable y alimentación solar. Cada escenario recalcula carbono, energía, coste aproximado y latencia, facilitando la elección de una arquitectura con menor impacto.")

doc.add_page_break()
doc.add_heading("6 Inteligencia artificial y aprendizaje automático", level=1)

doc.add_heading("6.1 Asistente de descarbonización", level=2)
add_body(doc, "El asistente recibe el proyecto activo, la evaluación calculada y la pregunta del usuario. FastAPI genera una respuesta contextual que prioriza medidas como filtrado por eventos en Edge, envío de telemetría por lotes, selección de regiones Cloud con baja intensidad de carbono y alimentación fotovoltaica de gateways.")
add_body(doc, "Si se configura una clave de un proveedor externo compatible, el backend puede solicitar un análisis generativo. Si no existe esa clave o la llamada falla, el sistema devuelve un diagnóstico local determinista. Por tanto, el módulo conserva una respuesta útil sin depender obligatoriamente de un servicio externo.")

doc.add_heading("6.2 Pipeline de aprendizaje automático", level=2)
add_body(doc, "El pipeline genera o carga un conjunto de datos, realiza análisis exploratorio, prepara variables, entrena modelos y evalúa su capacidad para estimar emisiones totales. La ejecución se inicia desde la interfaz y FastAPI informa el progreso hasta completar el reporte.")
add_table(doc, ["Modelo", "Papel en la evaluación"], [
    ["Random Forest", "Modelo de árboles robusto y base de comparación no lineal."],
    ["XGBoost", "Boosting de gradiente para relaciones complejas y alta precisión."],
    ["SVR", "Regresión de vectores de soporte como alternativa basada en márgenes."],
    ["Stacking", "Modelo híbrido que aprende a combinar predicciones de varios estimadores."],
    ["Blending", "Ensamble híbrido que pondera modelos a partir de un conjunto de validación."],
], [1.45, 5.35])

doc.add_heading("6.3 Evaluación estadística", level=2)
add_body(doc, "El reporte incluye RMSE, MAE, R cuadrado y MAPE, validación cruzada y pruebas estadísticas como Friedman, Wilcoxon, t de Nadeau y Bengio, bootstrap, Shapiro Wilk y Breusch Pagan. Estas pruebas permiten comparar modelos y revisar supuestos sobre los residuos. En la última ejecución verificada se entrenaron cinco modelos con 2500 registros y XGBoost fue seleccionado como mejor modelo por el criterio configurado.")

doc.add_page_break()
doc.add_heading("7 Persistencia seguridad y trazabilidad", level=1)

doc.add_heading("7.1 Persistencia PostgreSQL", level=2)
add_body(doc, "FastAPI expone operaciones de listado, alta o actualización y eliminación para proyectos, sensores, Edge, Cloud, gemelos, factores de emisión, escenarios, usuarios y auditorías. SQLAlchemy almacena cada recurso con su tipo, identificador externo, proyecto asociado y contenido JSONB.")

doc.add_heading("7.2 Control de acceso", level=2)
add_body(doc, "La interfaz distingue los roles Administrador, Investigador o Analista y Consulta. El rol activo determina qué acciones de edición se muestran o permiten. Para un despliegue productivo se recomienda complementar este control visual con autenticación, autorización en cada endpoint y gestión segura de sesiones.")

doc.add_heading("7.3 Auditoría", level=2)
add_body(doc, "El registro de auditoría conserva fecha, usuario, acción, entidad, detalle e información de origen cuando está disponible. Las exportaciones y modificaciones relevantes generan entradas que facilitan reconstruir qué se hizo y sobre qué recurso.")

doc.add_heading("7.4 Protección de secretos", level=2)
add_body(doc, "Las credenciales de PostgreSQL y las claves de proveedores de IA se configuran mediante variables de entorno en el backend. No deben incluirse en informes, capturas, repositorios públicos ni código fuente. En documentación compartida se deben utilizar marcadores como <contraseña> y <clave de API>.")

doc.add_page_break()
doc.add_heading("8 Flujo de uso recomendado", level=1)
add_numbered(doc, [
    "Crear o seleccionar el proyecto agrícola y revisar superficie, cultivo y periodo de análisis.",
    "Registrar sensores, nodos Edge y recursos Cloud con cantidades y consumos realistas.",
    "Definir los gemelos digitales, su frecuencia de sincronización y volumen de datos.",
    "Revisar y, si corresponde, actualizar los factores de emisión con su fuente técnica.",
    "Consultar el dashboard y el desglose del ciclo de vida para localizar los principales contribuyentes.",
    "Construir escenarios alternativos y comparar carbono, energía, coste y latencia.",
    "Usar el asistente IA como apoyo para formular medidas de reducción verificables.",
    "Ejecutar el pipeline ML cuando se necesite una estimación predictiva o comparar algoritmos.",
    "Exportar el informe y revisar el registro de auditoría antes de cerrar el análisis.",
])

doc.add_heading("8.1 Interpretación prudente", level=2)
add_body(doc, "Los resultados dependen de la calidad de los datos de inventario y de los factores de emisión. Las cifras deben entenderse como estimaciones técnicas del alcance modelado, no como una certificación ambiental automática. Para decisiones formales se recomienda documentar fuentes, año de referencia, límites del sistema y supuestos.")

doc.add_page_break()
doc.add_heading("9 Puesta en marcha sin Docker", level=1)
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(0)
spacer.paragraph_format.line_spacing = Pt(1)
add_body(doc, "El proyecto se ejecuta directamente sobre Windows, Node.js, Python y PostgreSQL. No requiere Docker.")

doc.add_heading("9.1 Frontend", level=2)
add_numbered(doc, [
    "Abrir PowerShell en la carpeta raíz del proyecto.",
    "Instalar dependencias con npm install.",
    "Iniciar Next.js con npm run dev.",
    "Abrir http://localhost:3000 en el navegador.",
])

doc.add_heading("9.2 Backend", level=2)
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(0)
spacer.paragraph_format.line_spacing = Pt(1)
add_numbered(doc, [
    "Entrar en la carpeta backend.",
    "Crear un entorno con python -m venv .venv.",
    "Activarlo con .venv\\Scripts\\Activate.ps1.",
    "Instalar dependencias con pip install -r requirements.txt.",
    "Configurar DATABASE_URL en backend/.env usando la contraseña real fuera del documento.",
    "Iniciar la API con python -m uvicorn app.main:app --reload --port 8000.",
])

doc.add_heading("9.3 Base de datos y servicios adicionales", level=2)
add_bullets(doc, [
    "PostgreSQL debe estar iniciado y la base indicada en DATABASE_URL debe existir.",
    "npm run db:seed comprueba y añade registros iniciales que falten sin sobrescribir los existentes.",
    "npm run ml:lab abre el laboratorio Streamlit en el puerto 8501 si están instaladas sus dependencias.",
    "La documentación interactiva de FastAPI queda disponible en http://localhost:8000/docs.",
])

doc.add_page_break()
doc.add_heading("10 Alcance actual y conclusiones", level=1)
add_body(doc, "Carbon Twin cubre el ciclo funcional completo de una herramienta de análisis: captura de inventario, persistencia, cálculo, visualización, simulación, recomendación, predicción y exportación. El visor 3D aporta una lectura espacial de la telemetría; el motor LCA convierte el inventario en indicadores ambientales; y los módulos de IA y ML permiten explorar medidas de mejora y patrones predictivos.")
add_body(doc, "La aplicación es adecuada como plataforma académica, prototipo avanzado y base para un producto de gestión ambiental agrícola. Antes de un uso productivo con múltiples organizaciones conviene añadir autenticación robusta, permisos aplicados en backend, migraciones de base de datos, pruebas automatizadas y monitorización.")

doc.add_heading("10.1 Resultado esperado", level=2)
add_body(doc, "Al finalizar un análisis, el usuario dispone de una huella total en kg CO2e, indicadores por hectárea y día, consumo energético, distribución por tecnología, desglose por etapa, comparación de escenarios, recomendaciones y archivos exportables. Todo queda vinculado al proyecto activo y respaldado por persistencia y auditoría.")

doc.add_heading("10.2 Criterios para considerar un análisis completo", level=2)
add_bullets(doc, [
    "El inventario incluye todos los equipos y recursos digitales relevantes.",
    "Los factores de emisión tienen unidad, fuente y periodo de referencia identificables.",
    "El horizonte temporal y los límites del sistema están definidos.",
    "Se ha revisado al menos un escenario alternativo de reducción.",
    "Las recomendaciones se han contrastado con restricciones técnicas y productivas.",
    "El informe exportado coincide con el proyecto y los datos visibles en el dashboard.",
])

# Propiedades y guardado
doc.core_properties.title = "Manual funcional del sistema Carbon Twin"
doc.core_properties.subject = "Análisis del coste de carbono de gemelos digitales agrícolas"
doc.core_properties.author = "AP 9 Carbon Twin"
doc.core_properties.keywords = "gemelo digital, carbono, agricultura, FastAPI, Next.js, PostgreSQL, LCA"
doc.save(OUTPUT)
print(OUTPUT)
