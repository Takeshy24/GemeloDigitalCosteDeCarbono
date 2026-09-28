from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / 'docs' / 'IEEE_ACCESS_ARTICLE.md'
OUT = ROOT / 'docs' / 'IEEE_ACCESS_ARTICLE.docx'

def set_cols(section, n=2, space='0.25'):
    sectPr = section._sectPr
    cols = sectPr.first_child_found_in('w:cols')
    if cols is None:
        cols = OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), str(int(float(space)*1440)))

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def clean(text):
    text = text.replace('**','').replace('*','')
    text = text.replace('--','-')
    text = re.sub(r'\$\$(.*?)\$\$', r'\1', text)
    text = re.sub(r'\$(.*?)\$', r'\1', text)
    return text.strip()

def add_para(doc, text, style='Body Text', size=9, bold=False, align=None, before=0, after=4):
    p = doc.add_paragraph(style=style if style in [s.name for s in doc.styles] else None)
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.0
    if align is not None: p.alignment=align
    r=p.add_run(clean(text)); r.bold=bold; r.font.name='Times New Roman'; r.font.size=Pt(size)
    return p

def add_table(doc, rows):
    table=doc.add_table(rows=0, cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            cell=cells[j]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ridx==0 else WD_ALIGN_PARAGRAPH.LEFT
            run=p.add_run(clean(val)); run.font.name='Times New Roman'; run.font.size=Pt(6.5 if len(row)>5 else 7.5); run.bold=(ridx==0)
            if ridx==0: set_cell_shading(cell,'1F4E79'); run.font.color.rgb=__import__('docx').shared.RGBColor(255,255,255)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

md=SRC.read_text(encoding='utf-8').splitlines()
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.0); sec.page_height=Inches(10.88); sec.top_margin=Inches(.89); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.51); sec.right_margin=Inches(.51)
set_cols(sec,1)
styles=doc.styles
styles['Normal'].font.name='Times New Roman'; styles['Normal'].font.size=Pt(9)

# Header/footer
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=header.add_run('IEEE ACCESS MANUSCRIPT'); rr.font.name='Arial'; rr.font.size=Pt(7); rr.bold=True
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rf=footer.add_run('CarbonTwin: Operational Carbon Estimation for Agricultural Digital Twins'); rf.font.name='Arial'; rf.font.size=Pt(7)

i=0; body_started=False
while i<len(md):
    line=md[i]
    if not line.strip(): i+=1; continue
    if line.startswith('# '):
        p=add_para(doc,line[2:],size=18,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,before=6,after=8)
        i+=1; continue
    if line.startswith('**Manuscript prepared'):
        add_para(doc,'',size=8,align=WD_ALIGN_PARAGRAPH.CENTER,after=2); i+=1; continue
    if line=='## Abstract':
        add_para(doc,'ABSTRACT',size=9,bold=True,align=WD_ALIGN_PARAGRAPH.LEFT,after=2); i+=1; continue
    if line.startswith('**Index Terms**'):
        add_para(doc,line.replace('**Index Terms**--','INDEX TERMS-'),size=8,bold=True,after=8)
        # begin two columns from first full section
        new=doc.add_section(WD_SECTION.CONTINUOUS)
        new.page_width=Inches(8.0); new.page_height=Inches(10.88); new.top_margin=Inches(.89); new.bottom_margin=Inches(.72); new.left_margin=Inches(.51); new.right_margin=Inches(.51); set_cols(new,2)
        body_started=True; i+=1; continue
    if line.startswith('## '):
        add_para(doc,line[3:].upper(),size=10,bold=True,before=7,after=3); i+=1; continue
    if line.startswith('**TABLE'):
        add_para(doc,line.replace('**',''),size=7.5,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,before=3,after=2); i+=1; continue
    if line.startswith('|') and i+1<len(md) and md[i+1].startswith('|---'):
        rows=[]
        while i<len(md) and md[i].startswith('|'):
            if not re.match(r'^\|[-:| ]+\|$',md[i]): rows.append([x.strip() for x in md[i].strip('|').split('|')])
            i+=1
        add_table(doc,rows); continue
    if re.match(r'^\d+\. ',line):
        p=add_para(doc,re.sub(r'^\d+\. ','',line),size=8.5,after=2); p.style='List Number'; i+=1; continue
    if line.startswith('$$'):
        eq=line.strip('$').replace('\\tag',' tag'); add_para(doc,eq,size=8,align=WD_ALIGN_PARAGRAPH.CENTER,after=4); i+=1; continue
    if line.startswith('[',) and re.match(r'^\[\d+\]',line):
        add_para(doc,line,size=7,after=1); i+=1; continue
    add_para(doc,line,size=8.5,after=4); i+=1

# Settings: refresh fields at Word open
settings=doc.settings.element
upd=OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true'); settings.append(upd)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
